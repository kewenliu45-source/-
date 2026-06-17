from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "智能库存预警系统_V1_客户确认版需求文档.md"
OUTPUT = ROOT / "docs" / "智能库存预警系统_V1_客户确认版需求文档.docx"


FONT_NAME = "Microsoft YaHei"
HEADING_BLUE = RGBColor(46, 116, 181)
HEADING_DARK_BLUE = RGBColor(31, 77, 120)
BODY_COLOR = RGBColor(31, 31, 31)


def set_run_font(run, font_name=FONT_NAME, size=None, bold=None, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY_COLOR
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, HEADING_BLUE, 16, 8),
        ("Heading 2", 13, HEADING_BLUE, 12, 6),
        ("Heading 3", 12, HEADING_DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def add_title_block(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, before=0, after=4, line=1.10)
    run = title.add_run("智能库存预警系统 V1.0")
    set_run_font(run, size=20, bold=True, color=RGBColor(11, 37, 69))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, before=0, after=18, line=1.10)
    run = subtitle.add_run("客户确认版需求文档")
    set_run_font(run, size=12, bold=False, color=RGBColor(85, 85, 85))

    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    widths = [Inches(1.4), Inches(5.1)]
    rows = [("文档用途", "用于客户确认第一阶段业务需求、功能范围与交付边界。"), ("确认范围", "登录、用户管理、低库存预警、缺码预警、高库存预警、调货建议、Excel 导出。")]
    for row, values in zip(table.rows, rows):
        for idx, text in enumerate(values):
            cell = row.cells[idx]
            cell.width = widths[idx]
            set_cell_shading(cell, "F2F4F7" if idx == 0 else "FFFFFF")
            paragraph = cell.paragraphs[0]
            set_paragraph_spacing(paragraph, before=0, after=0, line=1.10)
            run = paragraph.add_run(text)
            set_run_font(run, size=10.5, bold=(idx == 0), color=BODY_COLOR)

    doc.add_paragraph()


def add_markdown_line(doc, line):
    stripped = line.strip()
    if not stripped or stripped == "---":
        return
    if stripped.startswith("# "):
        if "客户确认版需求文档" not in stripped:
            paragraph = doc.add_heading(stripped[2:].strip(), level=1)
            for run in paragraph.runs:
                set_run_font(run, size=16, bold=True, color=HEADING_BLUE)
        return
    if stripped.startswith("## "):
        paragraph = doc.add_heading(stripped[3:].strip(), level=1)
        for run in paragraph.runs:
            set_run_font(run, size=16, bold=True, color=HEADING_BLUE)
        return
    if stripped.startswith("### "):
        paragraph = doc.add_heading(stripped[4:].strip(), level=2)
        for run in paragraph.runs:
            set_run_font(run, size=13, bold=True, color=HEADING_BLUE)
        return
    if stripped.startswith("#### "):
        paragraph = doc.add_heading(stripped[5:].strip(), level=3)
        for run in paragraph.runs:
            set_run_font(run, size=12, bold=True, color=HEADING_DARK_BLUE)
        return
    if stripped.startswith("- "):
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.167
        run = paragraph.add_run(stripped[2:].strip())
        set_run_font(run, size=11, color=BODY_COLOR)
        return

    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=6, line=1.10)
    if stripped.endswith("："):
        run = paragraph.add_run(stripped)
        set_run_font(run, size=11, bold=True, color=BODY_COLOR)
    else:
        run = paragraph.add_run(stripped)
        set_run_font(run, size=11, color=BODY_COLOR)


def add_footer(doc):
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("智能库存预警系统 V1.0 客户确认版")
    set_run_font(run, size=9, color=RGBColor(117, 117, 117))


def main():
    doc = Document()
    configure_styles(doc)
    add_title_block(doc)
    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        add_markdown_line(doc, line)
    add_footer(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
