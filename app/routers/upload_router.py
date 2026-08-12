import logging
import os
import shutil
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from io import BytesIO
from threading import Lock

from fastapi import APIRouter, UploadFile, File, Request
from fastapi.responses import HTMLResponse, StreamingResponse, FileResponse, JSONResponse
from fastapi.templating import Jinja2Templates

from app.auth import get_current_user
from app.config import BASE_DIR, DB_TYPE, OUTPUT_DIR

logger = logging.getLogger(__name__)
from app.data_sources.database_source import build_standard_data_from_database
from app.data_sources.excel_source import build_standard_data
from app.services.warning_service import analyze_standard_data, analyze_high_stock_data

router = APIRouter()
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "app", "templates"))

# 最近一次分析结果的文件路径（模块级变量）
_latest_result_file: str | None = None

# cpolar 等反向代理会中断长时间没有响应体的同步请求。
# 数据库分析因此使用短请求提交、后台执行和状态轮询。
_analysis_executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="inventory-analysis")
_analysis_tasks: dict[str, dict] = {}
_analysis_tasks_lock = Lock()
_ANALYSIS_TASK_TTL = timedelta(hours=6)
_analysis_uploads: dict[str, dict] = {}
_analysis_uploads_lock = Lock()
_ANALYSIS_UPLOAD_DIR = os.path.join(OUTPUT_DIR, "analysis_uploads")
_ANALYSIS_UPLOAD_TTL = timedelta(hours=2)
_ANALYSIS_UPLOAD_KINDS = {"hq_file", "transit_file", "inventory_file", "sales_file", "annual_sales_file"}

os.makedirs(_ANALYSIS_UPLOAD_DIR, exist_ok=True)


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _cleanup_analysis_tasks() -> None:
    cutoff = _utc_now() - _ANALYSIS_TASK_TTL
    with _analysis_tasks_lock:
        expired_ids = [
            task_id
            for task_id, task in _analysis_tasks.items()
            if task["created_at"] < cutoff
        ]
        for task_id in expired_ids:
            _analysis_tasks.pop(task_id, None)


def _create_analysis_task(owner: str) -> str:
    _cleanup_analysis_tasks()
    task_id = uuid.uuid4().hex
    now = _utc_now()
    with _analysis_tasks_lock:
        _analysis_tasks[task_id] = {
            "id": task_id,
            "owner": owner,
            "status": "queued",
            "message": "分析任务已提交，正在等待执行",
            "created_at": now,
            "updated_at": now,
            "result_context": None,
        }
    return task_id


def _update_analysis_task(task_id: str, **changes) -> None:
    with _analysis_tasks_lock:
        task = _analysis_tasks.get(task_id)
        if task is None:
            return
        task.update(changes)
        task["updated_at"] = _utc_now()


def _get_analysis_task(task_id: str, owner: str) -> dict | None:
    with _analysis_tasks_lock:
        task = _analysis_tasks.get(task_id)
        if task is None or task["owner"] != owner:
            return None
        return dict(task)


def _cleanup_analysis_uploads() -> None:
    cutoff = _utc_now() - _ANALYSIS_UPLOAD_TTL
    expired_dirs = []
    with _analysis_uploads_lock:
        expired_ids = [
            upload_id
            for upload_id, upload in _analysis_uploads.items()
            if upload["created_at"] < cutoff
        ]
        for upload_id in expired_ids:
            upload = _analysis_uploads.pop(upload_id, None)
            if upload:
                expired_dirs.append(upload["directory"])

    for directory in expired_dirs:
        shutil.rmtree(directory, ignore_errors=True)


def _create_analysis_upload(owner: str, files: dict) -> str:
    _cleanup_analysis_uploads()
    upload_id = uuid.uuid4().hex
    directory = os.path.join(_ANALYSIS_UPLOAD_DIR, upload_id)
    os.makedirs(directory, exist_ok=False)

    upload_files = {}
    for kind in _ANALYSIS_UPLOAD_KINDS:
        file_info = files.get(kind)
        if not file_info:
            continue
        size = int(file_info.get("size", 0))
        upload_files[kind] = {
            "name": str(file_info.get("name", "")),
            "size": size,
            "received": 0,
            "next_chunk": 0,
            "complete": size == 0,
            "path": os.path.join(directory, f"{kind}.xlsx"),
        }

    with _analysis_uploads_lock:
        _analysis_uploads[upload_id] = {
            "id": upload_id,
            "owner": owner,
            "directory": directory,
            "created_at": _utc_now(),
            "files": upload_files,
        }
    return upload_id


def _get_analysis_upload(upload_id: str, owner: str) -> dict | None:
    with _analysis_uploads_lock:
        upload = _analysis_uploads.get(upload_id)
        if upload is None or upload["owner"] != owner:
            return None
        return upload


def get_user_facing_error(exc: Exception) -> str:
    message = str(exc)
    if "EXERROR0002" in message and "invalid header line" in message:
        return (
            "畅捷通云端已转发到客户 T+ 账套服务，但 T+ 账套服务或中间代理返回了响应头格式非法的 HTTP 响应。"
            "错误中出现的 Date 响应头不是本系统发送的请求头。请在客户现场检查 T+ 服务、IIS/Nginx/网关代理"
            "是否改写了 Date 头，例如应为 `Date: Thu, 30 Jul 2026 11:25:39 GMT`，"
            "不能是 `Date Thu, 30 Jul 2026 11: 25:39 GMT`。修复代理/账套服务响应头后再重试。"
        )

    if "EXERROR0002" in message and "Connection refused" in message:
        return (
            "畅捷通已收到认证请求，但其内部无法连接当前 T+ 账套服务。"
            "请稍后重试；如果持续出现，请联系畅捷通技术支持，并提供错误码 "
            "EXERROR0002、组织 ID 和服务端日志中的 traceId。"
        )

    if "Read timed out" in message or ("接口在" in message and "秒内没有返回响应" in message):
        return (
            "调用畅捷通 T+ OpenAPI 超时。请先确认当前服务器能访问 openapi.chanjet.com，"
            "再根据现场网络和账套数据量调大 TPLUS_REQUEST_TIMEOUT，或调小 "
            "TPLUS_QUERY_PAGE_SIZE 后重试。"
        )

    return message


def build_result_context(result_df, data_source_name: str, warnings=None, output_filename=""):
    records = result_df.to_dict(orient="records")
    transfer_df = result_df[result_df["建议调货量"] > 0].copy()
    transfer_columns = [
        "存货编码",
        "存货",
        "尺码",
        "仓库编码",
        "仓库",
        "近7天销量",
        "日均销量",
        "当前可用量",
        "可售天数",
        "建议调货量",
        "总部可调数量",
        "在途仓",
        "在途（未发货）",
        "调货建议",
    ]
    transfer_columns = [column for column in transfer_columns if column in transfer_df.columns]
    transfer_records = transfer_df[transfer_columns].to_dict(orient="records")

    red_count = int((result_df["预警状态"] == "红色预警").sum())
    orange_count = int((result_df["预警状态"] == "橙色缺码").sum())
    yellow_count = int((result_df["预警状态"] == "黄色预警").sum())
    normal_count = int((result_df["预警状态"] == "正常").sum())

    suggest_total = int(result_df["建议调货量"].sum())
    hq_total = int(result_df["总部可调数量"].sum())

    summary = {
        "total": len(result_df),
        "red_count": red_count,
        "orange_count": orange_count,
        "yellow_count": yellow_count,
        "normal_count": normal_count,
        "suggest_total": suggest_total,
        "hq_total": hq_total,
        "data_source_name": data_source_name,
    }

    return {
        "records": records,
        "preview_data": records,
        "transfer_records": transfer_records,
        "summary": summary,
        "red_count": red_count,
        "orange_count": orange_count,
        "yellow_count": yellow_count,
        "normal_count": normal_count,
        "suggest_total": suggest_total,
        "hq_total": hq_total,
        "output_filename": output_filename,
        "warnings": warnings or [],
    }


def get_database_data_source_name() -> str:
    if DB_TYPE == "mock":
        return "模拟数据库"

    if DB_TYPE in {"tplus", "openapi", "chanjet"}:
        return "畅捷通 T+ OpenAPI"

    return "客户数据库"


def _save_result_excel(result_df, filepath: str, sheet_name: str = "预警结果", extra_colors: dict | None = None):
    """保存预警结果为带颜色的 Excel 文件。"""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    # 预警状态对应的颜色（行背景色）
    STATUS_COLORS = {
        "红色预警": "FCA5A5",    # 浅红
        "橙色缺码": "FDBA74",    # 浅橙
        "黄色预警": "FDE68A",    # 浅黄
        "正常":     "BBF7D0",    # 浅绿
        "高库存预警": "E9D5FF",  # 浅紫
    }
    if extra_colors:
        STATUS_COLORS.update(extra_colors)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 写表头
    columns = list(result_df.columns)
    header_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 找到预警状态列的索引
    status_col_idx = columns.index("预警状态") if "预警状态" in columns else -1

    # 写数据行
    for row_idx, row in enumerate(result_df.itertuples(index=False), 2):
        status = row[status_col_idx] if status_col_idx >= 0 else ""
        row_color = STATUS_COLORS.get(str(status).strip(), "")
        row_fill = PatternFill(start_color=row_color, end_color=row_color, fill_type="solid") if row_color else None

        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if row_fill:
                cell.fill = row_fill
            cell.alignment = Alignment(vertical="center")
            cell.border = thin_border

    # 表头也加边框
    for col_idx in range(1, len(columns) + 1):
        ws.cell(row=1, column=col_idx).border = thin_border

    # 自动调整列宽（取前100行的最大宽度）
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(str(col_name))
        for row_idx in range(2, min(102, ws.max_row + 1)):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val:
                max_len = max(max_len, len(str(cell_val)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 4, 40)

    # 冻结首行
    ws.freeze_panes = "A2"

    wb.save(filepath)


def prepare_analysis_result(standard_df, data_source_name: str, warnings=None):
    global _latest_result_file

    result = analyze_standard_data(standard_df)

    if isinstance(result, tuple):
        result_df = result[0]
    else:
        result_df = result

    # 保存分析结果为 Excel 文件，供下载（带预警颜色）
    try:
        filename = f"预警结果_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(OUTPUT_DIR, filename)
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        _save_result_excel(result_df, filepath)
        _latest_result_file = filepath
    except Exception as exc:
        logger.warning("保存预警结果文件失败: %s", exc)
        filename = ""

    context = build_result_context(
        result_df,
        data_source_name,
        warnings=warnings,
        output_filename=filename,
    )
    context["template_name"] = "result.html"
    return context


def render_analysis_result(request: Request, standard_df, data_source_name: str, warnings=None):
    return templates.TemplateResponse(
        request=request,
        name="result.html",
        context=prepare_analysis_result(standard_df, data_source_name, warnings=warnings),
    )


def _run_database_analysis_task(
    task_id: str,
    hq_file_bytes: bytes | None,
    transit_file_bytes: bytes | None,
) -> None:
    _update_analysis_task(
        task_id,
        status="running",
        message="正在从数据库读取并分析数据，请稍候",
    )

    try:
        from app.data_sources.excel_source import build_hq_standard_df, build_intransit_standard_df

        hq_df = None
        if hq_file_bytes:
            hq_df = build_hq_standard_df(BytesIO(hq_file_bytes))
            if hq_df.empty:
                hq_df = None

        transit_df = None
        if transit_file_bytes:
            transit_df = build_intransit_standard_df(BytesIO(transit_file_bytes))
            if transit_df.empty:
                transit_df = None

        standard_df = build_standard_data_from_database(hq_df=hq_df, transit_df=transit_df)
        result_context = prepare_analysis_result(
            standard_df,
            get_database_data_source_name(),
        )

        _update_analysis_task(
            task_id,
            status="succeeded",
            message="分析完成，正在打开结果页",
            result_context=result_context,
        )
    except Exception as exc:
        logger.exception("后台数据库分析任务失败: task_id=%s", task_id)
        _update_analysis_task(
            task_id,
            status="failed",
            message=get_user_facing_error(exc),
        )


def _run_excel_analysis_task(
    task_id: str,
    inventory_file_bytes: bytes,
    sales_file_bytes: bytes,
    transit_file_bytes: bytes | None,
    hq_file_bytes: bytes | None,
) -> None:
    _update_analysis_task(
        task_id,
        status="running",
        message="正在分析 Excel 数据，请稍候",
    )

    try:
        standard_df, warnings = build_standard_data(
            BytesIO(inventory_file_bytes),
            BytesIO(sales_file_bytes),
            BytesIO(transit_file_bytes) if transit_file_bytes else None,
            BytesIO(hq_file_bytes) if hq_file_bytes else None,
        )

        result_context = prepare_analysis_result(
            standard_df,
            "Excel 上传",
            warnings=warnings,
        )

        _update_analysis_task(
            task_id,
            status="succeeded",
            message="分析完成，正在打开结果页",
            result_context=result_context,
        )
    except Exception as exc:
        logger.exception("后台 Excel 分析任务失败: task_id=%s", task_id)
        _update_analysis_task(
            task_id,
            status="failed",
            message=get_user_facing_error(exc),
        )


def _run_high_stock_analysis_task(
    task_id: str,
    inventory_file_bytes: bytes,
    annual_sales_file_bytes: bytes,
) -> None:
    _update_analysis_task(
        task_id,
        status="running",
        message="正在分析高库存数据，请稍候",
    )

    try:
        from app.data_sources.excel_source import (
            build_inventory_standard_df,
            build_annual_sales_standard_df,
        )

        inventory_df = build_inventory_standard_df(BytesIO(inventory_file_bytes))
        annual_sales_df = build_annual_sales_standard_df(BytesIO(annual_sales_file_bytes))

        result_df = analyze_high_stock_data(inventory_df, annual_sales_df)

        # 保存结果 Excel
        output_filename = ""
        try:
            output_filename = f"高库存预警结果_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
            filepath = os.path.join(OUTPUT_DIR, output_filename)
            os.makedirs(OUTPUT_DIR, exist_ok=True)
            _save_result_excel(result_df, filepath, sheet_name="高库存预警结果")
        except Exception as exc:
            logger.warning("保存高库存预警结果文件失败: %s", exc)
            output_filename = ""

        result_context = build_high_stock_context(result_df, output_filename=output_filename)
        result_context["template_name"] = "high_stock_result.html"

        _update_analysis_task(
            task_id,
            status="succeeded",
            message="分析完成，正在打开结果页",
            result_context=result_context,
        )
    except Exception as exc:
        logger.exception("后台高库存分析任务失败: task_id=%s", task_id)
        _update_analysis_task(
            task_id,
            status="failed",
            message=get_user_facing_error(exc),
        )


def build_high_stock_context(result_df, warnings=None, output_filename=""):
    """构建高库存预警结果的模板 context。"""
    records = result_df.to_dict(orient="records")

    high_stock_count = int((result_df["预警状态"] == "高库存预警").sum())
    normal_count = int((result_df["预警状态"] == "正常").sum())
    high_stock_total_qty = int(
        result_df.loc[result_df["预警状态"] == "高库存预警", "当前可用量"].sum()
    )

    summary = {
        "total": len(result_df),
        "high_stock_count": high_stock_count,
        "normal_count": normal_count,
        "high_stock_total_qty": high_stock_total_qty,
        "data_source_name": "高库存 Excel 上传",
    }

    return {
        "records": records,
        "preview_data": records,
        "summary": summary,
        "output_filename": output_filename,
        "warnings": warnings or [],
    }


def render_high_stock_result(request: Request, result_df, warnings=None):
    """渲染高库存预警结果页面。"""
    global _latest_result_file

    try:
        filename = f"高库存预警结果_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(OUTPUT_DIR, filename)
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        _save_result_excel(result_df, filepath, sheet_name="高库存预警结果")
        _latest_result_file = filepath
    except Exception as exc:
        logger.warning("保存高库存预警结果文件失败: %s", exc)
        filename = ""

    return templates.TemplateResponse(
        request=request,
        name="high_stock_result.html",
        context=build_high_stock_context(result_df, warnings=warnings, output_filename=filename),
    )


@router.post("/high-stock-analysis", response_class=HTMLResponse)
async def analyze_high_stock(
    request: Request,
    inventory_file: UploadFile = File(...),
    annual_sales_file: UploadFile = File(...),
):
    """高库存预警分析入口。"""
    try:
        from app.data_sources.excel_source import (
            build_inventory_standard_df,
            build_annual_sales_standard_df,
        )

        inventory_df = build_inventory_standard_df(inventory_file.file)
        annual_sales_df = build_annual_sales_standard_df(annual_sales_file.file)

        result_df = analyze_high_stock_data(inventory_df, annual_sales_df)

        return render_high_stock_result(request, result_df)

    except Exception as exc:
        logger.exception("高库存分析请求处理失败")

        return templates.TemplateResponse(
            request=request,
            name="error.html",
            context={
                "error": get_user_facing_error(exc)
            }
        )


@router.get("/download/{filename}")
async def download_result_file(filename: str):
    """下载分析结果 Excel 文件"""
    filepath = os.path.join(OUTPUT_DIR, filename)
    if not os.path.exists(filepath):
        return HTMLResponse(content="文件不存在，请重新分析", status_code=404)
    return FileResponse(
        path=filepath,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
    )


@router.post("/upload", response_class=HTMLResponse)
async def upload_files(
    request: Request,
    inventory_file: UploadFile = File(...),
    sales_file: UploadFile = File(...),
    transit_file: UploadFile | None = File(None),
    hq_file: UploadFile = File(...),
):
    try:
        standard_df, warnings = build_standard_data(
            inventory_file.file,
            sales_file.file,
            transit_file.file if transit_file and transit_file.filename else None,
            hq_file.file if hq_file else None,
        )

        return render_analysis_result(request, standard_df, "Excel 上传", warnings=warnings)

    except Exception as exc:
        logger.exception("请求处理失败")

        return templates.TemplateResponse(
            request=request,
            name="error.html",
            context={
                "error": get_user_facing_error(exc)
            }
        )


@router.post("/analysis-uploads/init")
async def initialize_analysis_upload(request: Request):
    current_user = get_current_user(request) or {}
    payload = await request.json()
    files = payload.get("files", {}) if isinstance(payload, dict) else {}

    try:
        upload_id = _create_analysis_upload(
            current_user.get("username", ""),
            files,
        )
    except (TypeError, ValueError) as exc:
        return JSONResponse(
            {"error": f"上传文件信息无效: {exc}"},
            status_code=400,
        )

    return JSONResponse({"upload_id": upload_id})


@router.post("/analysis-uploads/{upload_id}/{kind}/chunks/{chunk_index}")
async def upload_analysis_chunk(
    request: Request,
    upload_id: str,
    kind: str,
    chunk_index: int,
):
    if kind not in _ANALYSIS_UPLOAD_KINDS:
        return JSONResponse({"error": "未知的上传文件类型"}, status_code=400)

    current_user = get_current_user(request) or {}
    owner = current_user.get("username", "")
    chunk = await request.body()
    if not chunk or len(chunk) > 1024 * 1024:
        return JSONResponse({"error": "分片大小无效"}, status_code=400)

    with _analysis_uploads_lock:
        upload = _analysis_uploads.get(upload_id)
        if upload is None or upload["owner"] != owner:
            return JSONResponse({"error": "上传任务不存在或已过期"}, status_code=404)

        file_info = upload["files"].get(kind)
        if file_info is None:
            return JSONResponse({"error": "该文件未登记"}, status_code=400)
        if chunk_index != file_info["next_chunk"]:
            return JSONResponse(
                {"error": "分片顺序不正确", "next_chunk": file_info["next_chunk"]},
                status_code=409,
            )
        if file_info["received"] + len(chunk) > file_info["size"]:
            return JSONResponse({"error": "上传数据超过文件大小"}, status_code=400)

        mode = "wb" if chunk_index == 0 else "ab"
        with open(file_info["path"], mode) as uploaded_file:
            uploaded_file.write(chunk)

        file_info["received"] += len(chunk)
        file_info["next_chunk"] += 1
        file_info["complete"] = file_info["received"] == file_info["size"]
        received = file_info["received"]
        complete = file_info["complete"]

    return JSONResponse({"received": received, "complete": complete})


@router.post("/database-analysis/start")
async def start_database_analysis(request: Request):
    current_user = get_current_user(request) or {}
    owner = current_user.get("username", "")
    payload = await request.json()
    upload_id = str(payload.get("upload_id", "")) if isinstance(payload, dict) else ""

    upload = _get_analysis_upload(upload_id, owner)
    if upload is None:
        return templates.TemplateResponse(
            request=request,
            name="error.html",
            context={"error": "上传任务不存在或已过期，请重新选择文件"},
            status_code=404,
        )

    incomplete = [
        file_info["name"]
        for file_info in upload["files"].values()
        if not file_info["complete"]
    ]
    if incomplete:
        return templates.TemplateResponse(
            request=request,
            name="error.html",
            context={"error": f"文件尚未上传完成: {', '.join(incomplete)}"},
            status_code=400,
        )

    hq_path = upload["files"].get("hq_file", {}).get("path")
    transit_path = upload["files"].get("transit_file", {}).get("path")
    if hq_path:
        with open(hq_path, "rb") as hq_upload:
            hq_file_bytes = hq_upload.read()
    else:
        hq_file_bytes = None
    if transit_path:
        with open(transit_path, "rb") as transit_upload:
            transit_file_bytes = transit_upload.read()
    else:
        transit_file_bytes = None

    with _analysis_uploads_lock:
        _analysis_uploads.pop(upload_id, None)
    shutil.rmtree(upload["directory"], ignore_errors=True)

    task_id = _create_analysis_task(owner)
    _analysis_executor.submit(
        _run_database_analysis_task,
        task_id,
        hq_file_bytes,
        transit_file_bytes,
    )

    return JSONResponse(
        {
            "task_id": task_id,
            "task_url": f"/analysis-tasks/{task_id}/result",
        },
        status_code=202,
    )


@router.post("/excel-analysis/start")
async def start_excel_analysis(request: Request):
    """启动 Excel 分析（低库存或高库存），支持分片上传后的后台执行。"""
    current_user = get_current_user(request) or {}
    owner = current_user.get("username", "")
    payload = await request.json()
    upload_id = str(payload.get("upload_id", "")) if isinstance(payload, dict) else ""
    kind = str(payload.get("kind", "low_stock")) if isinstance(payload, dict) else "low_stock"

    upload = _get_analysis_upload(upload_id, owner)
    if upload is None:
        return JSONResponse(
            {"error": "上传任务不存在或已过期，请重新选择文件"},
            status_code=404,
        )

    incomplete = [
        file_info["name"]
        for file_info in upload["files"].values()
        if not file_info["complete"]
    ]
    if incomplete:
        return JSONResponse(
            {"error": f"文件尚未上传完成: {', '.join(incomplete)}"},
            status_code=400,
        )

    def _read_upload_file(upload_files: dict, file_kind: str) -> bytes | None:
        info = upload_files.get(file_kind)
        if not info:
            return None
        with open(info["path"], "rb") as f:
            return f.read()

    files = upload["files"]

    # 先读取文件内容，再清理临时目录
    if kind == "high_stock":
        inventory_bytes = _read_upload_file(files, "inventory_file")
        annual_sales_bytes = _read_upload_file(files, "annual_sales_file")
    else:
        inventory_bytes = _read_upload_file(files, "inventory_file")
        sales_bytes = _read_upload_file(files, "sales_file")
        transit_bytes = _read_upload_file(files, "transit_file")
        hq_bytes = _read_upload_file(files, "hq_file")

    with _analysis_uploads_lock:
        _analysis_uploads.pop(upload_id, None)
    shutil.rmtree(upload["directory"], ignore_errors=True)

    task_id = _create_analysis_task(owner)

    if kind == "high_stock":
        _analysis_executor.submit(
            _run_high_stock_analysis_task,
            task_id,
            inventory_bytes,
            annual_sales_bytes,
        )
    else:
        _analysis_executor.submit(
            _run_excel_analysis_task,
            task_id,
            inventory_bytes,
            sales_bytes,
            transit_bytes,
            hq_bytes,
        )

    return JSONResponse(
        {
            "task_id": task_id,
            "task_url": f"/analysis-tasks/{task_id}/result",
        },
        status_code=202,
    )


@router.post("/database-analysis", response_class=HTMLResponse)
async def analyze_from_database(
    request: Request,
    hq_file: UploadFile | None = File(None),
    transit_file: UploadFile | None = File(None),
):
    current_user = get_current_user(request) or {}
    owner = current_user.get("username", "")

    hq_file_bytes = await hq_file.read() if hq_file and hq_file.filename else None
    transit_file_bytes = (
        await transit_file.read()
        if transit_file and transit_file.filename
        else None
    )

    task_id = _create_analysis_task(owner)
    _analysis_executor.submit(
        _run_database_analysis_task,
        task_id,
        hq_file_bytes,
        transit_file_bytes,
    )

    return templates.TemplateResponse(
        request=request,
        name="analysis_pending.html",
        context={"task_id": task_id},
        status_code=202,
    )


@router.get("/analysis-tasks/{task_id}")
async def get_analysis_task_status(request: Request, task_id: str):
    current_user = get_current_user(request) or {}
    task = _get_analysis_task(task_id, current_user.get("username", ""))
    if task is None:
        return JSONResponse(
            {"status": "not_found", "message": "分析任务不存在或已过期"},
            status_code=404,
        )

    response = {
        "status": task["status"],
        "message": task["message"],
    }
    if task["status"] == "succeeded":
        response["result_url"] = f"/analysis-tasks/{task_id}/result"
    return JSONResponse(response)


@router.get("/analysis-tasks/{task_id}/result", response_class=HTMLResponse)
async def get_analysis_task_result(request: Request, task_id: str):
    current_user = get_current_user(request) or {}
    task = _get_analysis_task(task_id, current_user.get("username", ""))
    if task is None:
        return templates.TemplateResponse(
            request=request,
            name="error.html",
            context={"error": "分析任务不存在或已过期，请重新分析"},
            status_code=404,
        )

    if task["status"] == "failed":
        return templates.TemplateResponse(
            request=request,
            name="error.html",
            context={"error": task["message"]},
            status_code=500,
        )

    if task["status"] != "succeeded":
        return templates.TemplateResponse(
            request=request,
            name="analysis_pending.html",
            context={"task_id": task_id},
            status_code=202,
        )

    template_name = task["result_context"].get("template_name", "result.html")
    return templates.TemplateResponse(
        request=request,
        name=template_name,
        context=task["result_context"],
    )


